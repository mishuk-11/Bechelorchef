"""Builds every STQA deliverable from the test design files and the real JUnit results.

    python docs/stqa/generate_reports.py --collect   # copy fresh JUnit XML from target/ into */evidence/, then build
    python docs/stqa/generate_reports.py             # build from the evidence already in the repo

Needs openpyxl, and Graphviz `dot` for the flow graphs (on PATH, in GRAPHVIZ_DOT, or in .devtools/graphviz).
Nothing about Pass/Fail is typed by hand: each row's status, time and failure text is read from the XML.
"""
from __future__ import annotations

import glob
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True  # design.py files are loaded as modules; keep the folders clean

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
MEMBERS = ["Inzamam", "Cherry", "Saima"]

NAVY = PatternFill("solid", fgColor="1F3864")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
GREY_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


# ------------------------------------------------------------------ helpers

def load_module(path: str, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def strip_output(src: str, dst: str) -> None:
    """Copy a JUnit XML file without the (large) captured stdout/stderr."""
    tree = ET.parse(src)
    for parent in tree.iter():
        for child in list(parent):
            if child.tag in ("system-out", "system-err"):
                parent.remove(child)
    tree.write(dst, encoding="utf-8", xml_declaration=True)


def collect() -> None:
    backend = os.path.join(ROOT, "backend", "target", "surefire-reports")
    e2e = os.path.join(ROOT, "e2e", "target", "surefire-reports")
    plan = {
        "Inzamam/evidence": [os.path.join(backend, "TEST-com.bachelorchef.stqa.inzamam.AdminManagement*.xml")],
        "Inzamam/evidence/unit": [os.path.join(backend, "TEST-com.bachelorchef.stqa.*.xml")],
        "Inzamam/evidence/selenium": [os.path.join(e2e, "TEST-com.bachelorchef.e2e.inzamam.*.xml")],
        "Cherry/evidence": [os.path.join(backend, "TEST-com.bachelorchef.stqa.cherry.*.xml")],
        "Saima/evidence": [os.path.join(backend, "TEST-com.bachelorchef.stqa.saima.*.xml")],
    }
    for target, patterns in plan.items():
        files = [f for p in patterns for f in glob.glob(p)]
        if not files:
            print(f"  ! no JUnit XML for {target} (run the tests first)")
            continue
        out = os.path.join(HERE, target)
        os.makedirs(out, exist_ok=True)
        for old in glob.glob(os.path.join(out, "*.xml")):
            os.remove(old)
        for f in files:
            strip_output(f, os.path.join(out, os.path.basename(f)))
        print(f"  collected {len(files)} file(s) -> docs/stqa/{target}")


def read_results(folder: str, id_pattern: str) -> dict:
    """Map test-case id (TC-07, SEL-12) -> result dict from every JUnit XML file in folder."""
    results = {}
    for path in sorted(glob.glob(os.path.join(folder, "*.xml"))):
        suite = ET.parse(path).getroot()
        stamp = (suite.get("timestamp") or "")[:10]
        for case in suite.iter("testcase"):
            m = re.search(id_pattern, case.get("name", ""))
            if not m:
                continue
            problem = case.find("failure")
            if problem is None:
                problem = case.find("error")
            skipped = case.find("skipped") is not None
            status = "Fail" if problem is not None else ("Skipped" if skipped else "Pass")
            message = ""
            if problem is not None:
                message = problem.get("message") or (problem.text or "").split("\n")[0]
            results[m.group(0)] = {
                "status": status,
                "ms": round(float(case.get("time", "0")) * 1000),
                "message": message,
                "date": stamp,
                "name": case.get("name"),
                "class": case.get("classname"),
            }
    return results


def short_failure(message: str) -> str:
    text = " ".join(message.split())
    m = re.search(r"Actual invocations? have different arguments(?: at position \[\d+\])?:\s*(.*)", text)
    if m:
        call = re.sub(r"\s*->\s*at .*", "", m.group(1))
        text = "service was called with " + call
    elif text.startswith("Wanted but not invoked"):
        text = "the service was never called (" + text[:120] + ")"
    text = text.replace("==> ", "").replace("org.opentest4j.AssertionFailedError: ", "")
    return text[:230] + ("..." if len(text) > 230 else "")


def actual_text(result: dict | None, defect: str) -> str:
    if result is None:
        return "Not executed (no JUnit result found)."
    if result["status"] == "Pass":
        return f"As expected. JUnit: Passed ({result['ms']} ms)."
    if result["status"] == "Skipped":
        return "Skipped (environment not configured)."
    who = f"Defect {defect}: " if defect else "Unexpected failure: "
    return f"{who}{short_failure(result['message'])}. JUnit: Failed ({result['ms']} ms)."


def style_header(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = NAVY
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = CENTER
        cell.border = BOX


def write_table(ws, start_row: int, headers: list, rows: list, widths: list, status_col: int | None = None) -> int:
    for i, h in enumerate(headers, 1):
        ws.cell(row=start_row, column=i, value=h)
    style_header(ws, start_row, len(headers))
    r = start_row
    for row in rows:
        r += 1
        for i, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.alignment = CENTER if (i == 1 or i == status_col) else WRAP
            cell.border = BOX
        if status_col:
            cell = ws.cell(row=r, column=status_col)
            cell.font = Font(bold=True, color="006100" if cell.value == "Pass" else "9C0006")
            if cell.value == "Pass":
                cell.fill = PASS_FILL
            elif cell.value == "Fail":
                cell.fill = FAIL_FILL
            else:
                cell.fill = GREY_FILL
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=2)
    return r


def title(ws, text: str, sub: str | None = None) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=13)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")


def md_escape(v) -> str:
    return str(v if v is not None else "").replace("|", "\\|").replace("\n", "<br>")


def md_table(headers: list, rows: list) -> str:
    out = ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"]
    out += ["| " + " | ".join(md_escape(c) for c in row) + " |" for row in rows]
    return "\n".join(out)


# ------------------------------------------------------------------ basis path

def graph_metrics(fn: dict) -> dict:
    nodes = len(fn["nodes"]) + 1  # + End
    edges = len(fn["edges"])
    p = sum(1 for n in fn["nodes"] if n[1] == "decision")
    return {"N": nodes, "E": edges, "P": p, "V_pred": p + 1, "V_graph": edges - nodes + 2}


def find_dot() -> str | None:
    for candidate in [os.environ.get("GRAPHVIZ_DOT"), shutil.which("dot"),
                      *glob.glob(os.path.join(ROOT, ".devtools", "graphviz", "*", "bin", "dot.exe"))]:
        if candidate and os.path.exists(candidate):
            return candidate
    return None


def draw_flow_graphs(design, out_png: str) -> bool:
    dot = find_dot()
    if not dot:
        print("  ! Graphviz not found; flow-graph PNG skipped")
        return False
    style = {
        "start": 'shape=box style="rounded,filled" fillcolor="#DDEBF7"',
        "process": 'shape=box style=filled fillcolor="#DDEBF7"',
        "decision": 'shape=diamond style=filled fillcolor="#FFF2CC"',
        "return": 'shape=box style="rounded,filled" fillcolor="#C6EFCE"',
        "error": 'shape=box style="rounded,filled" fillcolor="#F8CBAD"',
    }
    lines = ["digraph G {", 'graph [fontname="Arial" fontsize=12 labelloc=t newrank=true nodesep=0.35 ranksep=0.35',
             f'label="Control-flow graphs: {design.CONTROLLER} (owner {design.MEMBER}), 5 functions under basis path test"];',
             'node [fontname="Arial" fontsize=10 width=1.6];', 'edge [fontname="Arial" fontsize=10];']
    for k, fn in enumerate(design.FUNCTIONS):
        m = graph_metrics(fn)
        lines.append(f'subgraph cluster_{k} {{ style="rounded" color="#7F7F7F" fontsize=12 '
                     f'label="{fn["name"]}  (lines {fn["lines"]})\\nV(G) = P + 1 = {m["V_pred"]}   E - N + 2 = {m["E"]} - {m["N"]} + 2 = {m["V_graph"]}";')
        for nid, kind, line, text in fn["nodes"]:
            label = f"{nid}: {text}".replace('"', "'")
            wrapped = "\\n".join(_wrap(label, 30)) + f"\\n(line {line})"
            lines.append(f'  f{k}_{nid} [label="{wrapped}" {style[kind]}];')
        lines.append(f'  f{k}_End [label="End" shape=doublecircle width=0.4 fontsize=9 style=filled fillcolor="#D9D9D9"];')
        for a, b, lab in fn["edges"]:
            color = {"T": "#2E7D32", "F": "#C62828"}.get(lab, "#404040")
            attrs = f'color="{color}" fontcolor="{color}"' + (f' label=" {lab}"' if lab else "")
            lines.append(f"  f{k}_{a} -> f{k}_{b} [{attrs}];")
        lines.append("}")
    lines.append("}")
    dot_file = out_png[:-4] + ".dot"
    with open(dot_file, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    subprocess.run([dot, "-Tpng", "-Gdpi=110", dot_file, "-o", out_png], check=True)
    return True


def _wrap(text: str, width: int) -> list:
    words, out, cur = text.split(), [], ""
    for w in words:
        if cur and len(cur) + 1 + len(w) > width:
            out.append(cur)
            cur = w
        else:
            cur = f"{cur} {w}".strip()
    if cur:
        out.append(cur)
    return out


# ------------------------------------------------------------------ per member

def build_member(member: str, defects: dict) -> dict:
    folder = os.path.join(HERE, member)
    design = load_module(os.path.join(folder, "design.py"), f"design_{member}")
    results = read_results(os.path.join(folder, "evidence"), r"TC-\d\d")
    short = design.SHORT
    executed = sorted({r["date"] for r in results.values() if r["date"]}) or ["(not run)"]

    rows = []
    for tc, tech, fn, feature, data, steps, expected, method, defect in design.CASES:
        res = results.get(tc)
        status = res["status"] if res else "Not run"
        rows.append(dict(id=tc, technique=tech, function=fn, feature=feature, data=data, steps=steps, expected=expected,
                         method=method, defect=defect if status == "Fail" else "", planned_defect=defect,
                         actual=actual_text(res, defect), status=status, date=res["date"] if res else ""))
    counts = Counter(r["status"] for r in rows)
    by_tech = defaultdict(Counter)
    for r in rows:
        by_tech[r["technique"]][r["status"]] += 1

    # 1) Test cases workbook (the agreed 8-column format)
    wb = Workbook()
    ws = wb.active
    ws.title = "TestCases 01-30"
    title(ws, f"{design.CONTROLLER} - test cases TC-01..TC-30 (results from the JUnit run on {', '.join(executed)})",
          f"Owner: {member}. Test code: {design.TEST_DIR}. Status and Actual Result are copied from the JUnit XML in evidence/.")
    ws["A3"], ws["B3"], ws["C3"], ws["D3"], ws["E3"], ws["F3"] = "Total", f"=COUNTA(A6:A{5 + len(rows)})", "Pass", \
        f'=COUNTIF(G6:G{5 + len(rows)},"Pass")', "Fail", f'=COUNTIF(G6:G{5 + len(rows)},"Fail")'
    for c in ("A3", "C3", "E3"):
        ws[c].font = Font(bold=True)
    write_table(ws, 5, ["TC ID", "Action/Feature", "Test Data", "Steps", "Expected Result", "Actual Result", "Status", "xUnit Method"],
                [[r["id"], r["feature"], r["data"], r["steps"], r["expected"], r["actual"], r["status"], r["method"]] for r in rows],
                [8, 30, 30, 40, 36, 52, 9, 46], status_col=7)
    wb.save(os.path.join(folder, f"{short}_TestCases.xlsx"))

    # 2) Test design workbook: summary, BVA, ECT, full test-case sheet, defects found
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    title(ws, f"{design.CONTROLLER} - STQA test design and execution", f"Owner: {member}. Executed: {', '.join(executed)}.")
    summary_rows = [[t, sum(c.values()), c["Pass"], c["Fail"], c["Not run"] + c["Skipped"]] for t, c in by_tech.items()]
    summary_rows.append(["Total", len(rows), counts["Pass"], counts["Fail"], counts["Not run"] + counts["Skipped"]])
    write_table(ws, 4, ["Technique", "Cases", "Pass", "Fail", "Not run"], summary_rows, [16, 10, 10, 10, 10])

    ws = wb.create_sheet("BVA")
    title(ws, "Boundary Value Analysis", "Seven-point boundaries: min-1, min, min+1, nominal, max-1, max, max+1. '-' = not applicable.")
    write_table(ws, 4, ["Input", "Function", "Constraint (source)", "Below min", "Min", "Min+1", "Nominal", "Max-1", "Max", "Above max",
                        "Expected behaviour", "Test cases"], [list(b) for b in design.BVA], [24, 30, 34, 10, 8, 8, 9, 8, 8, 10, 36, 16])
    ws = wb.create_sheet("ECT")
    title(ws, "Equivalence Class Testing", "One representative value per valid and invalid class.")
    write_table(ws, 4, ["Input", "Class", "Description", "Valid / invalid", "Representative value", "Expected", "Test cases"],
                [list(e) for e in design.ECT], [24, 8, 30, 12, 30, 40, 18])
    ws = wb.create_sheet("TestCases")
    title(ws, "Test case design and execution (all 30 cases)")
    write_table(ws, 4, ["TC ID", "Function", "Technique", "Preconditions", "Test data", "Steps", "Expected result", "Automated test",
                        "Defect ID", "Actual result", "Status", "Executed on", "Executed by"],
                [[r["id"], r["function"], r["technique"], design.LOGIN, r["data"], r["steps"], r["expected"], r["method"],
                  r["defect"], r["actual"], r["status"], r["date"], member] for r in rows],
                [8, 16, 11, 24, 26, 34, 32, 40, 12, 50, 9, 12, 12], status_col=11)
    found = sorted({d.strip() for r in rows for d in r["defect"].split(",") if d.strip()})
    ws = wb.create_sheet("Defects")
    title(ws, f"Defects confirmed by {member}'s test run", "Full reports: Inzamam/BachelorChef_DefectReport.xlsx")
    write_table(ws, 4, ["Defect ID", "Title", "Severity", "Priority", "Location", "Confirmed by"],
                [[d, defects[d]["title"], defects[d]["severity"], defects[d]["priority"], defects[d]["location"],
                  ", ".join(r["id"] for r in rows if d in r["defect"])] for d in found], [10, 50, 10, 10, 46, 24])
    wb.save(os.path.join(folder, f"{short}_TestDesign.xlsx"))

    # 3) Basis path workbook + flow graph
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    title(ws, f"Basis path testing: {design.CONTROLLER}, 5 functions",
          "V(G) = decision nodes + 1, cross-checked with V(G) = E - N + 2 (End counted as a node). Source as of the test run.")
    srow, path_rows, legend_rows = [], [], []
    for fn in design.FUNCTIONS:
        m = graph_metrics(fn)
        pr = [results.get(p[4], {}).get("status", "Not run") for p in fn["paths"]]
        srow.append([fn["name"], fn["lines"], "; ".join(fn["predicates"]) or "(none)", m["P"], m["V_pred"], m["E"], m["N"],
                     m["V_graph"], "OK" if m["V_pred"] == m["V_graph"] else "MISMATCH", len(fn["paths"]),
                     "OK" if len(fn["paths"]) == m["V_pred"] else "CHECK", pr.count("Pass"), pr.count("Fail")])
        for pid, seq, cond, exp, tc in fn["paths"]:
            res = results.get(tc)
            path_rows.append([fn["name"], pid, " -> ".join(map(str, seq)) + " -> End", cond, exp, tc,
                              next(r["method"] for r in rows if r["id"] == tc), res["status"] if res else "Not run",
                              next(r["defect"] for r in rows if r["id"] == tc)])
        for nid, kind, line, text in fn["nodes"]:
            legend_rows.append([fn["name"], nid, kind, line, text])
    write_table(ws, 4, ["Function", "Source lines", "Decision predicates", "Decisions (P)", "V(G) = P + 1", "Edges (E)",
                        "Nodes (N, incl. End)", "V(G) = E - N + 2", "V(G) check", "Paths listed", "Paths = V(G)?", "Pass", "Fail"],
                srow, [16, 10, 46, 10, 10, 9, 10, 12, 10, 10, 11, 7, 7])
    ws = wb.create_sheet("Paths")
    title(ws, "Independent paths", "Node numbers match the 'Node Legend' sheet and the flow-graph PNG.")
    write_table(ws, 4, ["Function", "Path", "Node sequence", "Condition outcomes", "Expected result", "Test case", "JUnit method",
                        "Result", "Defect"], path_rows, [16, 6, 30, 34, 34, 9, 40, 9, 14], status_col=8)
    ws = wb.create_sheet("Node Legend")
    title(ws, "Node legend")
    write_table(ws, 3, ["Function", "Node", "Type", "Source line(s)", "Statement or decision"], legend_rows, [16, 6, 10, 12, 80])
    wb.save(os.path.join(folder, f"{short}_BasisPath.xlsx"))
    has_png = draw_flow_graphs(design, os.path.join(folder, f"{short}_FlowGraphs.png"))

    # 4) Markdown report
    md = [f"# {design.CONTROLLER} - STQA report ({member})", "",
          f"Source under test: [`{design.SOURCE}`](../../../{design.SOURCE})  ",
          f"Test code: [`{design.TEST_DIR}`](../../../{design.TEST_DIR})  ",
          f"Executed: {', '.join(executed)} with JUnit 5 + Mockito + MockMvc (standalone controller, real validation and error handler, mocked services).", "",
          "## Execution summary", "", md_table(["Technique", "Cases", "Pass", "Fail", "Not run"], summary_rows), "",
          "Every failing case confirms a seeded defect (see the Defect column); no test failed for any other reason." if
          all(r["defect"] or r["status"] != "Fail" for r in rows) else "Some failures are not linked to a defect: check them.", "",
          "## Test cases", "", md_table(["TC ID", "Action/Feature", "Test Data", "Expected Result", "Actual Result", "Status", "xUnit Method"],
                                         [[r["id"], r["feature"], r["data"], r["expected"], r["actual"], r["status"], r["method"]] for r in rows]), "",
          "## Boundary value analysis", "", md_table(["Input", "Function", "Constraint", "Min-1", "Min", "Min+1", "Nominal", "Max-1", "Max", "Max+1", "Expected", "Cases"],
                                                     [list(b) for b in design.BVA]), "",
          "## Equivalence classes", "", md_table(["Input", "Class", "Description", "Valid/invalid", "Representative", "Expected", "Cases"],
                                                  [list(e) for e in design.ECT]), "",
          "## Basis path testing", "",
          "V(G) is computed from the control-flow graph of the current source: decisions + 1, cross-checked with E - N + 2.", "",
          md_table(["Function", "Lines", "Decisions", "V(G)=P+1", "E", "N", "E-N+2", "Paths", "Pass", "Fail"],
                   [[s[0], s[1], s[3], s[4], s[5], s[6], s[7], s[9], s[11], s[12]] for s in srow]), "",
          (f"![Control-flow graphs]({short}_FlowGraphs.png)" if has_png else ""), "",
          md_table(["Function", "Path", "Node sequence", "Conditions", "Test case", "Result"],
                   [[p[0], p[1], p[2], p[3], p[5], p[7]] for p in path_rows]), "",
          "## Defects confirmed", "", md_table(["Defect", "Title", "Severity", "Confirmed by"],
                                               [[d, defects[d]["title"], defects[d]["severity"], ", ".join(r["id"] for r in rows if d in r["defect"])]
                                                for d in found]), "",
          "## Files", "",
          f"- `{short}_TestCases.xlsx` - the 30 test cases in the team format (TC ID ... xUnit Method)",
          f"- `{short}_TestDesign.xlsx` - summary, BVA table, ECT table, full test-case sheet, defects",
          f"- `{short}_BasisPath.xlsx` + `{short}_FlowGraphs.png` - V(G), independent paths, node legend, flow graphs",
          "- `evidence/` - JUnit XML of the run these files were generated from", ""]
    with open(os.path.join(folder, f"{short}_TestReport.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md))
    print(f"  {member}: {len(rows)} cases, {counts['Pass']} pass, {counts['Fail']} fail -> docs/stqa/{member}/")
    return {"member": member, "controller": design.CONTROLLER, "rows": rows, "counts": counts, "found": found}


# ------------------------------------------------------------------ Inzamam: defects, unit tests, Selenium

def build_selenium(defects: dict) -> dict:
    folder = os.path.join(HERE, "Inzamam")
    sel = load_module(os.path.join(folder, "selenium.py"), "selenium_cases")
    results = read_results(os.path.join(folder, "evidence", "selenium"), r"SEL-\d\d")
    rows = []
    for sid, feature, data, steps, expected, method, defect in sel.CASES:
        res = results.get(sid)
        status = res["status"] if res else "Not run"
        rows.append(dict(id=sid, feature=feature, data=data, steps=steps, expected=expected, method=method,
                         defect=defect if status == "Fail" else "", actual=actual_text(res, defect), status=status,
                         date=res["date"] if res else ""))
    executed = sorted({r["date"] for r in rows if r["date"]}) or ["(not run)"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Selenium"
    title(ws, f"Selenium UI automation - {len(rows)} cases (Chrome, results from the run on {', '.join(executed)})",
          f"Owner: Inzamam. Test code: {sel.TEST_DIR}. Needs the app running (run.bat).")
    write_table(ws, 4, ["TC ID", "Action/Feature", "Test Data", "Steps", "Expected Result", "Actual Result", "Status", "xUnit Method"],
                [[r["id"], r["feature"], r["data"], r["steps"], r["expected"], r["actual"], r["status"], r["method"]] for r in rows],
                [8, 30, 30, 44, 36, 52, 9, 42], status_col=7)
    wb.save(os.path.join(folder, "BachelorChef_SeleniumReport.xlsx"))
    counts = Counter(r["status"] for r in rows)
    print(f"  Selenium: {len(rows)} cases, {counts['Pass']} pass, {counts['Fail']} fail")
    return {"rows": rows, "counts": counts, "executed": executed}


def build_unit_report() -> dict:
    folder = os.path.join(HERE, "Inzamam", "evidence", "unit")
    per_class, all_cases = [], []
    for path in sorted(glob.glob(os.path.join(folder, "*.xml"))):
        suite = ET.parse(path).getroot()
        name = os.path.basename(path)[5:-4].replace("com.bachelorchef.", "")
        cases = list(suite.iter("testcase"))
        fails = [c for c in cases if c.find("failure") is not None or c.find("error") is not None]
        skips = [c for c in cases if c.find("skipped") is not None]
        owner = "Inzamam" if ".inzamam." in name else ("Cherry" if ".cherry." in name else "Saima")
        per_class.append([name, owner, len(cases), len(cases) - len(fails) - len(skips), len(fails), len(skips),
                          round(sum(float(c.get("time", 0)) for c in cases), 2)])
        for c in cases:
            status = "Fail" if c in fails else ("Skipped" if c in skips else "Pass")
            all_cases.append([name, c.get("name"), status])
    total = [sum(r[i] for r in per_class) for i in (2, 3, 4, 5)]
    wb = Workbook()
    ws = wb.active
    ws.title = "By class"
    title(ws, "JUnit 5 unit tests - whole backend (xUnit-family framework for Java)",
          f"{total[0]} tests: {total[1]} passed, {total[2]} failed, {total[3]} skipped. Failures are the seeded controller defects.")
    write_table(ws, 4, ["Test class", "Owner", "Tests", "Passed", "Failed", "Skipped", "Time (s)"], per_class, [62, 10, 8, 8, 8, 8, 9])
    ws = wb.create_sheet("All tests")
    title(ws, "Every executed test")
    write_table(ws, 3, ["Test class", "Test", "Status"], all_cases, [55, 90, 9], status_col=3)
    wb.save(os.path.join(HERE, "Inzamam", "BachelorChef_UnitTestReport.xlsx"))
    print(f"  Unit tests: {total[0]} tests, {total[1]} pass, {total[2]} fail")
    return {"per_class": per_class, "total": total}


def build_defect_report(defects: dict, member_results: list, selenium: dict) -> None:
    found_by = defaultdict(list)
    for mr in member_results:
        for r in mr["rows"]:
            for d in [x.strip() for x in r["defect"].split(",") if x.strip()]:
                found_by[d].append(f"{r['id']} ({mr['member']})")
    for r in selenium["rows"]:
        for d in [x.strip() for x in r["defect"].split(",") if x.strip()]:
            found_by[d].append(f"{r['id']} (Selenium)")
    owner = {**{f"BUG-{i:02d}": ("AdminManagementController", "Inzamam") for i in range(1, 11)},
             **{f"BUG-{i:02d}": ("ProfileController", "Cherry") for i in range(11, 21)},
             **{f"BUG-{i:02d}": ("PersonalController", "Saima") for i in range(21, 31)}}
    rows = []
    for d in defects.values():
        module, tester = owner[d["id"]]
        steps = d["steps"] + (("\nNote: " + d["note"]) if d["note"] else "")
        rows.append([d["id"], d["title"], module, d["severity"], d["priority"], d["technique"], d["location"], steps,
                     d["expected"], d["actual"], d["cause"], ", ".join(found_by.get(d["id"], [])) or "Manual only",
                     "Yes" if found_by.get(d["id"]) else "No", tester, "Open"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Defects"
    sev = Counter(r[3] for r in rows)
    title(ws, f"Bachelor Chef - defect report ({len(rows)} defects)",
          f"Reported by Inzamam. Severity: " + ", ".join(f"{k} {v}" for k, v in sev.items()) +
          ". Every defect is confirmed by at least one failing automated test unless 'Confirmed by test' says No.")
    write_table(ws, 4, ["Defect ID", "Title", "Module", "Severity", "Priority", "Technique", "Location", "Steps to reproduce",
                        "Expected result", "Actual result", "Root cause", "Confirmed by (test case)", "Confirmed by test",
                        "Module tester", "Status"], rows, [9, 40, 22, 9, 9, 14, 30, 50, 36, 40, 34, 30, 10, 11, 8])
    ws = wb.create_sheet("By module")
    by_mod = Counter((r[2], r[3]) for r in rows)
    write_table(ws, 1, ["Module", "Minor", "Low", "Total"],
                [[m, by_mod[(m, "Minor")], by_mod[(m, "Low")], by_mod[(m, "Minor")] + by_mod[(m, "Low")]]
                 for m in ("AdminManagementController", "ProfileController", "PersonalController")], [30, 8, 8, 8])
    wb.save(os.path.join(HERE, "Inzamam", "BachelorChef_DefectReport.xlsx"))
    md = ["# Bachelor Chef - defect report", "",
          f"Reported by **Inzamam**. {len(rows)} defects, all Open. Severity: " + ", ".join(f"{k} {v}" for k, v in sev.items()) + ".",
          "Each defect lists the automated test cases that confirm it (team unit tests and Selenium).", "",
          md_table(["ID", "Title", "Module", "Severity", "Priority", "Confirmed by"], [[r[0], r[1], r[2], r[3], r[4], r[11]] for r in rows]), ""]
    for r in rows:
        md += [f"## {r[0]} - {r[1]}", "",
               f"- **Module:** {r[2]} (tested by {r[13]})  ",
               f"- **Severity / priority:** {r[3]} / {r[4]}  ",
               f"- **Location:** `{r[6]}`  ",
               f"- **Confirmed by:** {r[11]}", "",
               "**Steps to reproduce**", "", "```", r[7], "```", "",
               f"**Expected:** {r[8]}  ", f"**Actual:** {r[9]}  ", f"**Root cause:** {r[10]}", ""]
    with open(os.path.join(HERE, "Inzamam", "BachelorChef_DefectReport.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md).replace("\n**Expected:**", "\n**Expected:**"))
    print(f"  Defects: {len(rows)} ({sum(1 for r in rows if r[12] == 'Yes')} confirmed by automated tests)")


def build_overview(member_results: list, selenium: dict, unit: dict) -> None:
    lines = ["# STQA results overview", "", "Generated by `generate_reports.py` from the JUnit XML in each `evidence/` folder.", "",
             md_table(["Member", "Module", "Cases", "Pass", "Fail", "Defects confirmed"],
                      [[m["member"], m["controller"], len(m["rows"]), m["counts"]["Pass"], m["counts"]["Fail"], len(m["found"])]
                       for m in member_results]), "",
             f"Selenium (Inzamam): {len(selenium['rows'])} cases, {selenium['counts']['Pass']} pass, {selenium['counts']['Fail']} fail.  ",
             f"JUnit unit tests, whole backend (Inzamam): {unit['total'][0]} tests, {unit['total'][1]} pass, {unit['total'][2]} fail.", ""]
    with open(os.path.join(HERE, "RESULTS.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def main() -> None:
    if "--collect" in sys.argv:
        print("Collecting JUnit XML ...")
        collect()
    print("Building reports ...")
    defects = {d["id"]: d for d in json.load(open(os.path.join(HERE, "Inzamam", "defects.json"), encoding="utf-8"))}
    member_results = [build_member(m, defects) for m in MEMBERS]
    selenium = build_selenium(defects)
    unit = build_unit_report()
    build_defect_report(defects, member_results, selenium)
    build_overview(member_results, selenium, unit)
    print("Done.")


if __name__ == "__main__":
    main()
